#!/usr/bin/env python3
"""Build candidate-v0.2 model inputs for the 12 pilot papers (docs/master_table/pilot_v0.1.csv).

T1: new structured instruction and output schema (evaluation/schemas/output_v0.2.json).
    Material is reused from the frozen v0.1 T1 input when one exists, otherwise from the
    v0.1 T3 manuscript_before (same sanitised text and page images), otherwise built from
    the PeerJ submission PDF with the v0.1 PeerJ convention (author cover pages dropped,
    text from the Abstract page on, every later page imaged).
T2: v0.1 inputs copied byte for byte (format unchanged in v0.2).
T3: v0.1 inputs with the v0.2 schema (issue_type added). The two deferred items get the
    supplements recovered in intake batch-01. PeerJ cycles v0.1 -> v0.2 are built new from
    the submission PDFs, the batch-01 first-round reports and the response letter.
Requires the candidate-v0.1.0 assets and the batch-01 Release unpacked at --intake-root.
No models are called.
"""
import argparse
import csv
import copy
import hashlib
import json
import re
import shutil
from pathlib import Path

import docx
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'pilot/inputs'
SCHEMAS = json.loads((ROOT / 'evaluation/schemas/output_v0.2.json').read_text())
CATALOG = json.loads((ROOT / 'evaluation/source_catalog.json').read_text())
DPI = 160
PEERJ_CROP_TOP = 36
PEERJ_CROP_BOTTOM = 40

ISSUE_TYPES = ('V1 construct and measurement validity; V2 internal validity and causal inference; '
               'V3 statistical conclusion validity (analysis choice, numerical consistency, power); '
               'V4 external validity and generalisation; V5 qualitative trustworthiness (credibility, '
               'transferability, dependability, confirmability); V6 theory and argument; '
               'V7 reporting transparency and reproducibility; V8 conclusions proportionate to evidence; '
               'V9 non-substantive (language, formatting, citation style)')

T1_INSTRUCTION = (
    'You are reviewing the supplied social-science manuscript using only the provided material. '
    'Identify up to five of the most important substantive scientific issues, ordered by importance. '
    'For each issue state the claim, classify it with one issue_type, say whether it is a confirmed_error, '
    'missing_information or suggestion, rate its severity (major or minor), give the reasoning (why it is a '
    'problem and which conclusion it affects) and support it with evidence: an exact quote copied verbatim '
    'from manuscript_text, or a page number from page_images with what you observed there. '
    f'Issue types: {ISSUE_TYPES}. '
    'Also give a summary (at most 150 words) of the central claims and an overall_assessment (at most 150 words). '
    'Do not invent missing analyses or references. Do not infer facts from the paper\'s publication status. '
    'Use the supplied page images to read tables, figures, equations and any ambiguous text-extraction order; '
    'manuscript_text is a searchable aid and must not replace visual material. External resources mentioned in '
    'the manuscript are not available for this task. Their non-inclusion does not establish that the authors failed '
    'to share them. If a judgment requires such resources, state the limitation. Do not browse, open links or '
    'retrieve other manuscript versions. Return only JSON that validates against output_schema.')

T3_ISSUE_NOTE = f'\n每个问题须用issue_type标注一个类型：{ISSUE_TYPES}。'

T3_PEERJ_INSTRUCTION = (
    '你是社会科学研究的独立评议者。阅读以下同一稿件一个修订周期的全部给定材料，判断修订稿在当前证据下的实质可接受性。'
    '给定的报告、编辑决定信和作者回复是待核查的主张，不是金标准；你不需要服从审稿人的错误要求，也不能把作者说“已修改”当作已解决。\n'
    '本题来自期刊发表前的评审；判断不等于编辑的接收决定或下一轮结果。时点为作者提交修订稿后、下一轮评审之前。'
    '只根据本题的材料和附带图表判断，不检索原文、不利用文献身份或历史结局。\n'
    '先列出最多5个最影响核心主张的已解决或残留问题，再给总体判断。细微语言/排版问题本身不足以判“需要实质修改”。'
    '若缺少某项外部数据，请明确证据边界；不要仅因本题没有原始数据就自动判证据不足。\n'
    '每个重要判断都要引用材料字段名和准确短引文/页图位置；区分回应主张、可见实际修改、科学上是否足够。所有字段按以下JSON结构输出。'
    + T3_ISSUE_NOTE)

MATERIAL_SCOPE_PEERJ = ('完整当轮主文、参考书目及嵌入图表。作者封面页不提供；正文自摘要页起，其后各页均提供页面图像。'
                        '未附作者在正文中提到的外部原始数据/代码/仓储资源；不得据此推断作者未共享。')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')


def copy_asset(src, dest_dir, name):
    dest = dest_dir / name
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dest)
    return dest


def peerj_pdf(pdf, dest_dir, prefix):
    """Sanitised text and page images of a PeerJ submission PDF, following the v0.1 convention."""
    doc = pymupdf.open(pdf)
    # Crop the PeerJ running header (logo, 'Manuscript to be reviewed') and footer ('PeerJ reviewing PDF | id').
    clip = lambda page: pymupdf.Rect(0, PEERJ_CROP_TOP, page.rect.width, page.rect.height - PEERJ_CROP_BOTTOM)
    start = next(i for i, page in enumerate(doc)
                 if re.search(r'^\s*Abstract\b', page.get_text(clip=clip(page)), re.M | re.I))
    text = ''.join(doc[i].get_text(clip=clip(doc[i])) for i in range(start, len(doc)))
    pages = []
    for i in range(start + 1, len(doc)):
        name = f'{prefix}_p{i + 1:03d}.png'
        target = dest_dir / name
        target.parent.mkdir(parents=True, exist_ok=True)
        doc[i].get_pixmap(dpi=DPI, clip=clip(doc[i])).save(target)
        pages.append({'path': name, 'page_label': f'p{i + 1:03d}'})
    return text, pages, {'source_pdf': str(Path(pdf).relative_to(ROOT)), 'source_sha256': sha(pdf),
                         'text_from_pdf_page': start + 1, 'crop_points': [PEERJ_CROP_TOP, PEERJ_CROP_BOTTOM], 'imaged_pages': [p['page_label'] for p in pages], 'dpi': DPI}


def docx_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text for cell in row.cells))
    return '\n'.join(parts)


SUPPLEMENT_FOOTER = re.compile(r'doi\.org|\bDOI:?\s*10\.\d{4}|CC BY 4\.0\s*\|')
PEERJ_SECTIONS = ('Basic reporting', 'Experimental design', 'Validity of the findings', 'Additional comments',
                  'Comments for the Author', 'Comments for the author')


def clean_peerj_review(text):
    """Drop the reviewer-name/permalink header and the 'Cite this review as' block (title and DOI)."""
    starts = [text.find(h) for h in PEERJ_SECTIONS if h in text]
    start = min(starts) if starts else 0
    end = text.find('Cite this review as')
    return text[start:end if end != -1 else None].strip()


def clean_response_letter(text, title):
    """Drop identifying text: salutation naming the editor, an author-name heading, reviewer names in
    headings, the paper title, and a closing signature (only near the end, since some letters put a
    signed cover note before the point-by-point replies)."""
    text = re.sub(r'\A\s*Dear [^\n]*\n', '', text)
    text = re.sub(r'\A\s*Response to Reviewers?\b[^\n]*\n', 'Response to Reviewers\n', text, flags=re.I)
    text = re.sub(r'(Reviewer\s*#?\s*\d+)\s*\([^)\n]*\)', r'\1', text)
    text = re.sub(re.escape(title.strip()), '[title removed]', text, flags=re.I)
    closing = r'^[ \t]*(Kind regards|Best regards|Sincerely|Yours sincerely|Regards)\b[^\n]*\n'
    tail = len(text) * 0.85
    ends = [m for m in re.finditer(closing, text, re.M | re.I) if m.start() >= tail]
    if ends:
        text = text[:ends[-1].start()]
    # A signed cover note before the replies: drop the closing line and the short name lines under it.
    text = re.sub(closing + r'(?:[ \t]*[^\n]{1,120}\n){0,3}', '', text, flags=re.M | re.I)
    return text.strip()


def pdf_supplement(pdf, dest_dir, key):
    """Text and page images of a supplement, with DOI/attribution footer lines redacted in both."""
    doc = pymupdf.open(pdf)
    text, pages = [], []
    for page in doc:
        for block in page.get_text('dict')['blocks']:
            block_text = ' '.join(span['text'] for line in block.get('lines', []) for span in line['spans'])
            if SUPPLEMENT_FOOTER.search(block_text) and len(block_text) < 200:
                page.add_redact_annot(block['bbox'], fill=(1, 1, 1))
        page.apply_redactions()
    for i, page in enumerate(doc):
        text.append(f'\n[Source attachment page {i + 1}]\n' + page.get_text())
        name = f'assets/supplement_{key}/page_{i + 1:03d}.png'
        target = dest_dir / name
        target.parent.mkdir(parents=True, exist_ok=True)
        page.get_pixmap(dpi=DPI).save(target)
        pages.append({'path': name, 'source_attachment_page': i + 1})
    return ''.join(text), pages


def t1_from_v01(task_dir, src_input):
    """T1 v0.2 from an existing T1 v0.1 input; returns (input, lineage)."""
    old = json.loads(src_input.read_text())
    new_in = copy.deepcopy(old['input'])
    for img in new_in.get('page_images', []):
        copy_asset(src_input.parent / img['file'], task_dir, img['file'])
    for s in new_in.get('supplements', {}).values():
        for vp in s.get('visual_pages', []):
            path = vp.get('path') or vp.get('file')
            copy_asset(src_input.parent / path, task_dir, path)
    return new_in, {'material_from': str(src_input.relative_to(ROOT)), 'material_sha256': sha(src_input)}


def t1_from_t3(task_dir, t3_input):
    old = json.loads(t3_input.read_text())
    before = old['materials']['manuscript_before']
    pages = []
    for p in before.get('visual_pages', []):
        name = 'materials/pages/' + Path(p['path']).name.replace('manuscript_before_', 'page_')
        copy_asset(t3_input.parent / p['path'], task_dir, name)
        pages.append({'file': name, 'page_label': p.get('page_label') or f"p{p.get('source_page', 0):03d}"})
    supplements = {}
    for key, s in old['materials'].get('supplements', {}).items():
        if 'manuscript_before' not in s.get('applies_to', []):
            continue
        s = copy.deepcopy(s)
        for vp in s.get('visual_pages', []):
            copy_asset(t3_input.parent / vp['path'], task_dir, vp['path'])
        supplements[key] = s
    new_in = {'manuscript_text': before['text'], 'page_images': pages,
              'material_scope': '完整当轮主文、参考书目及嵌入图表，与该论文 T3 题 manuscript_before 的净化材料相同。'
                                '未附作者在正文中提到的外部原始数据/代码/仓储资源；不得据此推断作者未共享。'}
    if supplements:
        new_in['supplements'] = supplements
    return new_in, {'material_from': str(t3_input.relative_to(ROOT)) + '#/materials/manuscript_before',
                    'material_sha256': sha(t3_input)}


def t1_from_peerj(task_dir, ledger):
    text, pages, src = peerj_pdf(ROOT / ledger['前稿原件'], task_dir, 'materials/pages/page')
    pages = [{'file': p['path'], 'page_label': p['page_label']} for p in pages]
    return {'manuscript_text': text, 'page_images': pages, 'material_scope': MATERIAL_SCOPE_PEERJ}, src


def build_t1(sample, pid, ledger, cards):
    task_id = f'{sample}_T1_Q01'
    task_dir = OUT / task_id
    if task_dir.exists():
        shutil.rmtree(task_dir)
    if ('T1', task_id) in cards:
        inp, lineage = t1_from_v01(task_dir, ROOT / cards[('T1', task_id)])
    elif ('T3', f'T3_{sample}_R1') in cards:
        inp, lineage = t1_from_t3(task_dir, ROOT / cards[('T3', f'T3_{sample}_R1')])
    else:
        inp, lineage = t1_from_peerj(task_dir, ledger)
    obj = {'task_id': task_id, 'paper_id': pid, 'task': 'T1', 'unit': 'original_manuscript', 'version': '0.2',
           'required_modalities': ['text', 'image'] if inp.get('page_images') else ['text'],
           'instruction': T1_INSTRUCTION, 'input': inp, 'output_schema': SCHEMAS['schemas']['T1']}
    dump(task_dir / 'input.json', obj)
    return task_id, lineage


def build_t2(cards):
    built = []
    for (task, task_id), path in sorted(cards.items()):
        if task != 'T2':
            continue
        src = ROOT / path
        task_dir = OUT / task_id
        if task_dir.exists():
            shutil.rmtree(task_dir)
        task_dir.mkdir(parents=True)
        shutil.copyfile(src, task_dir / 'input.json')
        built.append((task_id, {'copied_from': path, 'sha256': sha(src), 'changes': 'none (T2 format unchanged in v0.2)'}))
    return built


def t3_schema_upgrade(obj):
    obj['output_schema'] = SCHEMAS['schemas']['T3']
    obj['instruction'] = obj['instruction'].rstrip() + T3_ISSUE_NOTE
    obj['version'] = '0.2'
    obj['required_modalities'] = ['text', 'image'] if '"visual_pages"' in json.dumps(obj) else ['text']
    return obj


def build_t3_existing(task_id, src, intake_fill):
    task_dir = OUT / task_id
    if task_dir.exists():
        shutil.rmtree(task_dir)
    shutil.copytree(src.parent, task_dir, ignore=shutil.ignore_patterns('private', '评分规则.txt', '题面.txt', '暂缓原因.txt'))
    obj = json.loads(src.read_text())
    lineage = {'copied_from': str(src.relative_to(ROOT)), 'sha256': sha(src), 'changes': ['v0.2 output schema with issue_type']}
    if (src.parent / '暂缓原因.txt').is_file():
        sup = obj['materials'].setdefault('supplements', {})
        meta = json.loads((intake_fill / 'metadata.json').read_text())
        for n, entry in enumerate(meta['supplements'], start=len(sup) + 1):
            key = f'A{n:02d}'
            text, pages = pdf_supplement(intake_fill / entry['file'], task_dir, key)
            sup[key] = {'text': text, 'visual_pages': pages, 'applies_to': ['manuscript_after'],
                        'description': f"作者在后稿中指向的外部正式附件；DOI 与署名页脚已遮盖；"
                                       f"许可 {entry.get('license')}；来源 intake batch-01"}
            lineage['changes'].append(f"supplement {key} from intake batch-01 {entry['file']} sha256={entry['sha256']}")
        obj['supplement_policy'] = (obj.get('supplement_policy', '') +
                                    ' 原暂缓题所缺的外部附件已由 intake batch-01 补入，见 applies_to=manuscript_after 的条目。').strip()
    dump(task_dir / 'input.json', t3_schema_upgrade(obj))
    return lineage


def build_t3_peerj(sample, pid, ledger, intake_fill):
    task_id = f'T3_{sample}_R1'
    task_dir = OUT / task_id
    if task_dir.exists():
        shutil.rmtree(task_dir)
    before_text, before_pages, before_src = peerj_pdf(ROOT / ledger['前稿原件'], task_dir, 'assets/manuscript_before')
    after_text, after_pages, after_src = peerj_pdf(ROOT / ledger['修改稿原件'], task_dir, 'assets/manuscript_after')
    meta = json.loads((intake_fill / 'metadata.json').read_text())
    reviews = {r['id'].split('_')[0]: clean_peerj_review((intake_fill / r['text_file']).read_text(encoding='utf-8'))
               for r in meta['reviews'] if r['on_version'] == 'v1'}
    responses, response_src = {}, []
    for r in meta['responses']:
        path = ROOT / r['file'] if r.get('existing_repository_file') else intake_fill / r['file']
        responses['combined'] = clean_response_letter(docx_text(path), ledger['题名'])
        response_src.append({'file': str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else r['file'],
                             'sha256': sha(path), 'answers': r['answers'], 'date': r['date']})
    version_dir = (ROOT / ledger['修改稿原件']).parent
    decision = sorted(version_dir.glob('*_decision.txt'))
    editor = re.sub(r'^Editor:.*\n', '', decision[0].read_text(encoding='utf-8')) if decision else None
    materials = {'manuscript_before': {'text': before_text, 'visual_pages': before_pages, 'version_role': 'before'},
                 'reviews': reviews, 'responses': responses,
                 'manuscript_after': {'text': after_text, 'visual_pages': after_pages, 'version_role': 'after'}}
    if editor:
        materials['editor_decision'] = editor
    obj = {'question_id': task_id, 'task': 'T3', 'paper_id': pid, 'version': '0.2',
           'required_modalities': ['text', 'image'], 'instruction': T3_PEERJ_INSTRUCTION,
           'timepoint': '作者提交修订稿后、下一轮评审之前；仅含针对前稿的第一轮报告、编辑决定信与作者答复',
           'process': 'pre_publication_journal_revision', 'materials': materials,
           'output_schema': SCHEMAS['schemas']['T3'],
           'supplement_policy': '未提供外部底层原始数据；不因未附推断作者未共享。'}
    dump(task_dir / 'input.json', obj)
    return task_id, {'built_new': True, 'manuscript_before': before_src, 'manuscript_after': after_src,
                     'reviews': [{'id': r['id'], 'text_file': f"intake/batch-01/{intake_fill.name}/{r['text_file']}",
                                  'date': r['date']} for r in meta['reviews'] if r['on_version'] == 'v1'],
                     'responses': response_src, 'editor_decision': str(decision[0].relative_to(ROOT)) if decision else None,
                     'editor_name_line_removed': bool(editor),
                     'review_cleaning': 'reviewer name/permalink header and citation block (title, DOI) removed',
                     'response_cleaning': 'salutation, reviewer names in headings and signature removed'}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--intake-root', type=Path, required=True,
                        help='checkout of intake/batch-01 with its Release archives unpacked')
    args = parser.parse_args()
    pilot = [r for r in csv.DictReader((ROOT / 'docs/master_table/pilot_v0.1.csv').open(encoding='utf-8'))
             if r['pilot_role'] != 'unused']
    all_cards = {}
    for f in (ROOT / 'evaluation/cards').glob('*.json'):
        c = json.loads(f.read_text())
        all_cards.setdefault(c['paper_id'], {})[(c['task'], c['task_id'])] = c['input_path']
    manifest = []
    for row in pilot:
        pid = row['paper_id']
        ledger = CATALOG[pid]['ledger']
        sample = ledger['样本编号']
        cards = all_cards.get(pid, {})
        fill = args.intake_root / 'intake/batch-01' / ('fill_' + pid.replace(':', '_'))
        entries = []
        task_id, lineage = build_t1(sample, pid, ledger, cards)
        entries.append((task_id, 'T1', lineage))
        for task_id, lineage in build_t2(cards):
            entries.append((task_id, 'T2', lineage))
        t3 = cards.get(('T3', f'T3_{sample}_R1'))
        if t3:
            entries.append((f'T3_{sample}_R1', 'T3', build_t3_existing(f'T3_{sample}_R1', ROOT / t3, fill)))
        elif pid.startswith('PeerJ:'):
            task_id, lineage = build_t3_peerj(sample, pid, ledger, fill)
            entries.append((task_id, 'T3', lineage))
        for task_id, task, lineage in entries:
            path = OUT / task_id / 'input.json'
            manifest.append({'task_id': task_id, 'paper_id': pid, 'pilot_role': row['pilot_role'], 'task': task,
                             'input_path': str(path.relative_to(ROOT)), 'input_sha256': sha(path), 'lineage': lineage})
    for m in manifest:
        task_dir = OUT / m['task_id']
        m['assets'] = {str(f.relative_to(task_dir)): sha(f) for f in sorted(task_dir.rglob('*.png'))}
    dump(OUT / 'manifest.json', {'version': '0.2-pilot', 'models_called': False, 'pymupdf': pymupdf.VersionBind,
                                 'png_policy': 'PNG files are not committed; rebuild with this script and check assets hashes',
                                 'items': manifest})
    counts = {}
    for m in manifest:
        counts[(m['pilot_role'], m['task'])] = counts.get((m['pilot_role'], m['task']), 0) + 1
    print(len(manifest), 'inputs;', {f'{k[0]}/{k[1]}': v for k, v in sorted(counts.items())})


if __name__ == '__main__':
    main()
