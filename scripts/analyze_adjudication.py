#!/usr/bin/env python3
"""Unblind the filled first-adjudication workbook and summarise how each drafter's points fared.

Rebuilds the row order of scripts/build_adjudication_workbook.py from pilot/merge/, checks it
against the filled workbook row by row, then maps X/Y choices to drafters A (Claude) / B (GPT)
with pilot/blinding_key.json. Writes pilot/unblinded/adjudication_results/first_adjudication.json and a
per-row CSV. Usage: python3 scripts/analyze_adjudication.py [filled.xlsx]
"""
import csv
import json
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MERGE = ROOT / 'pilot/merge'
OUT = ROOT / 'pilot/unblinded/adjudication_results'  # git-ignored until the second adjudication is done
DEFAULT = ROOT / 'pilot/adjudication_filled/裁决表_盲评版_第一裁决.xlsx'
ORDER = {'T2': 0, 'T1': 1, 'T3': 2}


def expected_rows():
    merged = [json.loads(p.read_text(encoding='utf-8')) for p in MERGE.glob('*_*.json')]
    merged.sort(key=lambda m: (ORDER[m['X']['task']], m['task_id']))
    for m in merged:
        xs = {p['id']: p for p in m['X']['reference_points']}
        ys = {p['id']: p for p in m['Y']['reference_points']}
        for group in m['suggested_alignment']:
            px = [xs[i] for i in group['x'].split(';') if i]
            py = [ys[i] for i in group['y'].split(';') if i]
            kind = 'both' if px and py else ('X_only' if px else 'Y_only')
            yield {'task_id': m['task_id'], 'task': m['X']['task'], 'kind': kind, 'X': px, 'Y': py}
        if m['X']['task'] != 'T1':
            yield {'task_id': m['task_id'], 'task': m['X']['task'], 'kind': 'label',
                   'X': m['X']['label_reference'], 'Y': m['Y']['label_reference']}


def sign_test(k, n):
    """Two-sided exact binomial test against p = 0.5."""
    tail = sum(math.comb(n, i) for i in range(0, min(k, n - k) + 1)) / 2 ** n
    return min(1.0, 2 * tail)


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    key = json.loads((ROOT / 'pilot/blinding_key.json').read_text())['tasks']
    filled = [r for r in load_workbook(path)['裁决表'].iter_rows(min_row=2, values_only=True) if r[0]]
    rows = list(expected_rows())
    assert len(rows) == len(filled), (len(rows), len(filled))
    records = []
    for exp, got in zip(rows, filled):
        assert exp['task_id'] == got[0], (exp['task_id'], got[0])
        side = {'X': key[exp['task_id']]['X'], 'Y': key[exp['task_id']]['Y']}
        choice = got[4]
        rec = {'task_id': exp['task_id'], 'task': exp['task'], 'kind': exp['kind'], 'choice': choice, 'note': got[5] or ''}
        if exp['kind'] == 'label':
            rec['agree'] = exp['X']['primary'] == exp['Y']['primary']
            rec['winner'] = {'X对': side['X'], 'Y对': side['Y'], '都对': 'both', '都不对': 'neither'}[choice]
        elif exp['kind'] == 'both':
            rec['winner'] = {'X对': side['X'], 'Y对': side['Y'], '都对': 'both', '都不对': 'neither'}[choice]
            rec['card_author'] = {side[s]: any(p['origin'] == 'card_author' for p in exp[s]) for s in 'XY'}
        else:
            owner_side = exp['kind'][0]
            rec['owner'] = side[owner_side]
            rec['accepted'] = choice in (f'{owner_side}对', '都对')
            rec['card_author'] = any(p['origin'] == 'card_author' for p in exp[owner_side])
            rec['importance'] = sorted({p['importance'] for p in exp[owner_side]})
        records.append(rec)

    summary = {'source': str(path.relative_to(ROOT)), 'rows': len(records),
               'choices': Counter(r['choice'] for r in records), 'time_recorded': any(r[6] for r in filled)}
    both = [r for r in records if r['kind'] == 'both']
    summary['both_sides'] = {t: dict(Counter(r['winner'] for r in both if r['task'] == t)) for t in ('T1', 'T2', 'T3')}
    summary['both_sides']['all'] = dict(Counter(r['winner'] for r in both))
    a, b = summary['both_sides']['all'].get('A', 0), summary['both_sides']['all'].get('B', 0)
    summary['both_sides']['sign_test_A_vs_B'] = {'A': a, 'B': b, 'p_two_sided': round(sign_test(a, a + b), 4)}
    one = [r for r in records if r['kind'] in ('X_only', 'Y_only')]
    acc = defaultdict(Counter)
    for r in one:
        acc[f"{r['owner']}_{'card_author' if r['card_author'] else 'human_review'}"]['accepted' if r['accepted'] else 'rejected'] += 1
        acc[f"{r['owner']}_{r['task']}"]['accepted' if r['accepted'] else 'rejected'] += 1
    summary['one_side'] = {k: dict(v) for k, v in sorted(acc.items())}
    labels = [r for r in records if r['kind'] == 'label']
    summary['labels'] = {'drafters_agreed': dict(Counter(r['winner'] for r in labels if r['agree'])),
                         'drafters_disagreed': dict(Counter(r['winner'] for r in labels if not r['agree']))}
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'first_adjudication.json').write_text(json.dumps(summary, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    with (OUT / 'first_adjudication_rows.csv').open('w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['task_id', 'task', 'kind', 'choice', 'winner', 'owner', 'accepted', 'card_author', 'importance', 'agree', 'note'])
        w.writeheader()
        for r in records:
            w.writerow({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v) for k, v in r.items()})
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
