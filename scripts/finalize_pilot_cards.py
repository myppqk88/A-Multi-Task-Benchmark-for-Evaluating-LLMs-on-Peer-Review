#!/usr/bin/env python3
"""Turn the first-adjudication workbook into pilot reference cards (evaluation/cards_pilot/).

Rules (recorded in evaluation/测量框架与效度论证.md §6.3):
- Two-sided row, "X对"/"Y对": keep that side's points; the other side's are dropped.
- Two-sided row, "都对": keep the side with fewer points (coarser grain, one point per issue; ties
  keep X, which was assigned at random) and add the other side's judgments as acceptable
  alternatives; quoted evidence from both sides is pooled.
- One-sided row accepted: keep the point. Rejected ("都不对"): the adjudicator did not say whether it
  was wrong or not worth raising, so it goes to neutral_points, which count neither towards recall
  nor as invalid critique.
- Label row: the chosen side's label; "都对" takes the union of the acceptable sets, and the primary
  is the label both drafters accepted, taking the more critical one if two remain.
Provenance names sides as X/Y only, so the cards do not reveal which drafter wrote what.
Usage: python3 scripts/finalize_pilot_cards.py [filled.xlsx]
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from analyze_adjudication import DEFAULT, expected_rows  # noqa: E402

OUT = ROOT / 'evaluation/cards_pilot'
SECOND = set(json.loads((ROOT / 'pilot/second_adjudication_sample.json').read_text())['second_adjudication'])
CRITICAL_FIRST = ['insufficient', 'partially_sufficient', 'sufficient', 'needs_substantive_revision',
                  'acceptable_under_rubric', 'insufficient_evidence']
# The same issue drafted by both sides but left in separate rows by the evidence-overlap grouping.
# (task, kept point, duplicate folded into it as an acceptable alternative), checked by hand.
DUPLICATES = [('D16_T1_Q01', 'X:P7', 'Y:P16'), ('D20_T1_Q01', 'Y:P8', 'X:P10')]


def fold_duplicates(task, points):
    for t, keep, drop in DUPLICATES:
        if t != task:
            continue
        k = next(p for p in points if p['provenance']['from'] == keep)
        d = next(p for p in points if p['provenance']['from'] == drop)
        k['acceptable_alternatives'] = k.get('acceptable_alternatives', []) + [d['judgment']]
        k['human_sources'] = pooled([k, d], 'human_sources')
        k['manuscript_evidence'] = pooled([k, d], 'manuscript_evidence')
        k['provenance']['folded_duplicate'] = drop
        points.remove(d)


def pooled(points, field):
    seen, out = set(), []
    for p in points:
        for item in p.get(field, []):
            k = json.dumps(item, sort_keys=True, ensure_ascii=False)
            if k not in seen:
                seen.add(k)
                out.append(item)
    return out


def final_label(row, choice):
    if choice in ('X对', 'Y对'):
        lab = dict(row[choice[0]])
        lab['rule'] = f'adjudicator chose {choice[0]}'
        return lab
    x, y = row['X'], row['Y']
    both = [l for l in CRITICAL_FIRST if l in x['acceptable_set'] and l in y['acceptable_set']]
    primary = x['primary'] if x['primary'] == y['primary'] else (both[0] if both else x['primary'])
    union = [l for l in CRITICAL_FIRST if l in set(x['acceptable_set']) | set(y['acceptable_set'])]
    return {'primary': primary, 'acceptable_set': union,
            'rationale': {'X': x.get('rationale', ''), 'Y': y.get('rationale', '')},
            'rule': 'both accepted: union of sets; primary agreed by both drafters, more critical first'}


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    filled = [r for r in load_workbook(path)['裁决表'].iter_rows(min_row=2, values_only=True) if r[0]]
    rows = list(expected_rows())
    assert len(rows) == len(filled)
    merged = {p.stem: json.loads(p.read_text(encoding='utf-8')) for p in (ROOT / 'pilot/merge').glob('*_*.json')}
    cards = {}
    for n, (row, got) in enumerate(zip(rows, filled), 1):
        assert row['task_id'] == got[0]
        task, choice = row['task_id'], got[4]
        m = merged[task]
        card = cards.setdefault(task, {
            'task_id': task, 'paper_id': m['X']['paper_id'], 'task': m['X']['task'], 'unit': m['X']['unit'],
            'input_sha256': m['X']['input_sha256'], 'card_version': 'pilot-1.0',
            'status': 'provisional_pending_second_adjudication' if task in SECOND else 'final_first_adjudication',
            'reference_points': [], 'neutral_points': [], 'label_reference': None,
            'known_invalid_patterns': {'X': m['X'].get('known_invalid_patterns', []), 'Y': m['Y'].get('known_invalid_patterns', [])},
            'known_invalid_patterns_note': '起草者提出、未经裁决，仅供评分者参考。'})
        if row['kind'] == 'label':
            card['label_reference'] = final_label(row, choice)
            continue
        if row['kind'] == 'both':
            if choice == '都对':
                keep = 'X' if len(row['X']) <= len(row['Y']) else 'Y'
            else:
                keep = choice[0]
            other = 'Y' if keep == 'X' else 'X'
            for p in row[keep]:
                q = {k: v for k, v in p.items() if k != 'id'}
                if choice == '都对':
                    q['acceptable_alternatives'] = p.get('acceptable_alternatives', []) + [o['judgment'] for o in row[other]]
                    q['human_sources'] = pooled([p] + row[other], 'human_sources')
                    q['manuscript_evidence'] = pooled([p] + row[other], 'manuscript_evidence')
                q['provenance'] = {'sheet_row': n + 1, 'from': f"{keep}:{p['id']}", 'decision': choice}
                card['reference_points'].append(q)
        else:
            side = row['kind'][0]
            accepted = choice in (f'{side}对', '都对')
            for p in row[side]:
                q = {k: v for k, v in p.items() if k != 'id'}
                q['provenance'] = {'sheet_row': n + 1, 'from': f"{side}:{p['id']}", 'decision': choice}
                card['reference_points' if accepted else 'neutral_points'].append(q)
    OUT.mkdir(parents=True, exist_ok=True)
    totals = {'tasks': 0, 'reference_points': 0, 'high': 0, 'neutral_points': 0}
    for task, card in sorted(cards.items()):
        fold_duplicates(task, card['reference_points'])
        for i, p in enumerate(card['reference_points'], 1):
            card['reference_points'][i - 1] = {'id': f'R{i:02d}', **p}
        for i, p in enumerate(card['neutral_points'], 1):
            card['neutral_points'][i - 1] = {'id': f'N{i:02d}', **p}
        card['recall_denominator_high'] = sum(p['importance'] == 'high' and p['status'] == 'supported' for p in card['reference_points'])
        (OUT / f'{task}.json').write_text(json.dumps(card, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
        totals['tasks'] += 1
        totals['reference_points'] += len(card['reference_points'])
        totals['high'] += card['recall_denominator_high']
        totals['neutral_points'] += len(card['neutral_points'])
    print(json.dumps(totals, ensure_ascii=False))


if __name__ == '__main__':
    main()
