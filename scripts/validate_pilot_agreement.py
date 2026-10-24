#!/usr/bin/env python3
"""Two agreement checks for the pilot (测量框架 §7, 评分环节).

1. Human L1 scoring vs the LLM judge on the same claims (pilot/adjudication_filled/补评表_模型答案_已填.xlsx):
   agreement on the decision category and, where both say "match", on the point matched.
2. Second vs first adjudicator on the 8 overlap tasks (裁决表_第二裁决_已填.xlsx vs the first workbook).
Cohen's kappa is unweighted. Writes pilot/scores/agreement.json.
"""
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILLED = ROOT / 'pilot/adjudication_filled'
HUMAN = {'对上参照点': 'match', '新问题·成立': 'valid_new', '成立但不重要': 'valid_minor', '不成立': 'invalid', '无法判断': 'unverifiable'}


def kappa(pairs):
    n = len(pairs)
    if not n:
        return None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return round((po - pe) / (1 - pe), 3) if pe < 1 else None


def human_vs_judge():
    key = json.loads((ROOT / 'pilot/unblinded/补评表_key.json').read_text())['tasks']
    sheet = {r: v for r, v in enumerate(load_workbook(FILLED / '补评表_模型答案_已填.xlsx')['补评表'].iter_rows(min_row=1, values_only=True), 1)}
    pairs, point_pairs, rows = [], [], []
    for task_id, info in key.items():
        judge = json.loads((ROOT / 'pilot/scores/l1_judge' / f'{task_id}.json').read_text())
        inverse = {m: l for l, m in judge['label_key'].items()}
        for r in info['rows']:
            human = sheet[r['sheet_row']]
            h = HUMAN.get(human[4])
            decisions = judge['judgement']['answers'][inverse[r['model']]]
            d = next((x for x in decisions if x['claim_field'] == r['claim_field'] and x['claim_index'] == r['claim_index']), None)
            if h is None or d is None:
                continue
            j = 'match' if d['decision'] == 'neutral' else d['decision']
            j = 'valid_minor' if d['decision'] == 'neutral' else j
            pairs.append((h, j))
            if h == 'match' and j == 'match':
                point_pairs.append(((human[5] or '').strip().upper(), d.get('point_id')))
            rows.append({'task_id': task_id, 'model': r['model'], 'human': h, 'judge': d['decision'],
                         'human_point': human[5], 'judge_point': d.get('point_id'), 'human_note': human[6]})
    by_model = {}
    for m in ('opus', 'sonnet', 'haiku'):
        mp = [(x['human'], 'valid_minor' if x['judge'] == 'neutral' else x['judge']) for x in rows if x['model'] == m]
        by_model[m] = {'n': len(mp), 'agree': round(sum(a == b for a, b in mp) / len(mp), 3) if mp else None}
    return {'n': len(pairs), 'percent_agreement': round(sum(a == b for a, b in pairs) / len(pairs), 3), 'kappa': kappa(pairs),
            'human_counts': Counter(a for a, _ in pairs), 'judge_counts': Counter(b for _, b in pairs),
            'confusion': Counter(f'{a}->{b}' for a, b in pairs if a != b),
            'same_point_when_both_match': f'{sum(a == b for a, b in point_pairs)}/{len(point_pairs)}',
            'by_model': by_model, 'rows': rows}


def first_vs_second():
    first = [r for r in load_workbook(FILLED / '裁决表_盲评版_第一裁决.xlsx')['裁决表'].iter_rows(min_row=2, values_only=True) if r[0]]
    second = [r for r in load_workbook(FILLED / '裁决表_第二裁决_已填.xlsx')['裁决表'].iter_rows(min_row=2, values_only=True) if r[0]]
    index = {}
    for r in first:
        index.setdefault((r[0], r[2], r[3]), []).append(r[4])
    pairs, unmatched = [], 0
    for r in second:
        hit = index.get((r[0], r[2], r[3]))
        if not hit:
            unmatched += 1
            continue
        pairs.append((hit.pop(0), r[4]))
    by_kind = {}
    for kind, test in (('two_sided', lambda p: True),):
        by_kind[kind] = None
    return {'n': len(pairs), 'unmatched_rows': unmatched, 'percent_agreement': round(sum(a == b for a, b in pairs) / len(pairs), 3),
            'kappa': kappa(pairs), 'first_counts': Counter(a for a, _ in pairs), 'second_counts': Counter(b for _, b in pairs),
            'confusion': Counter(f'{a}->{b}' for a, b in pairs if a != b),
            'accept_reject_agreement': round(sum((a == '都不对') == (b == '都不对') for a, b in pairs) / len(pairs), 3)}


def main():
    out = {'human_vs_judge': human_vs_judge(), 'first_vs_second_adjudicator': first_vs_second()}
    (ROOT / 'pilot/scores/agreement.json').write_text(json.dumps(out, ensure_ascii=False, indent=1, default=dict) + '\n', encoding='utf-8')
    for k, v in out.items():
        print(k, {x: y for x, y in v.items() if x != 'rows'})


if __name__ == '__main__':
    main()
