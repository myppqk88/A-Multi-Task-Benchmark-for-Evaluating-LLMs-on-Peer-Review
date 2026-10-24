#!/usr/bin/env python3
"""Blind human L1 scoring sheet for a seeded sample of pilot tasks, to validate the LLM judge.

Samples tasks (seed 20261028; T1 and T2 only, excluding the second-adjudication tasks whose cards
are provisional), lists every claim of every model answer for those tasks under fresh anonymous
labels, and asks for the same decision the judge makes. Neither the model names nor the judge's
decisions appear in the sheet. The label key goes to git-ignored pilot/unblinded/.
Writes pilot/补评表_模型答案.xlsx and pilot/unblinded/补评表_key.json.
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
RUNS = ROOT / 'pilot/runs'
CARDS = ROOT / 'evaluation/cards_pilot'
MODELS = ('opus', 'sonnet', 'haiku')
SEED = 20261028
PER_TASK = {'T1': 2, 'T2': 2}
CHOICES = ['对上参照点', '新问题·成立', '成立但不重要', '不成立', '无法判断']
IMP = {'high': '高', 'medium': '中', 'low': '低'}
FONT, BOLD = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
FILL = PatternFill('solid', start_color='FFFF99')
HEAD = PatternFill('solid', start_color='D9E1F2')
EDGE = Side(style='thin', color='BBBBBB')
BORDER = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def claims(answer, task):
    if task == 'T2':
        for field, prefix in (('remaining_gaps', '仍存在的缺口'), ('addressed_requirements', '已处理的要求')):
            for i, text in enumerate(answer.get(field, [])):
                yield field, i, f'〔{prefix}〕{text}'
    else:
        for i, issue in enumerate(answer.get('issues', [])):
            text = issue.get('claim') or issue.get('issue', '')
            reason = issue.get('reasoning') or issue.get('observed_change', '')
            yield 'issues', i, f'{text}\n——理由：{reason}'


def cell(ws, row, col, value, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = BOLD if bold else FONT
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if fill:
        c.fill = fill
    return c


def main():
    second = set(json.loads((ROOT / 'pilot/second_adjudication_sample.json').read_text())['second_adjudication'])
    manifest = json.loads((ROOT / 'pilot/inputs/manifest.json').read_text())['items']
    rng = random.Random(SEED)
    chosen = []
    for task, n in PER_TASK.items():
        pool = sorted(m['task_id'] for m in manifest if m['task'] == task and m['task_id'] not in second
                      and all((RUNS / mod / f"{m['task_id']}.json").is_file() for mod in MODELS))
        chosen += rng.sample(pool, n)
    chosen.sort(key=lambda t: (0 if '_T2_' in t else 1, t))

    wb = Workbook()
    guide = wb.active
    guide.title = '说明'
    lines = [
        ('这张表要你做什么', True),
        ('几个 AI 模型分别审了同一道题。每一行是某个模型提出的一条意见。请你判断这条意见属于哪一类。', False),
        ('模型名字已经隐藏，用"答案甲/乙/丙"代替；同一道题里，甲乙丙是不同的模型。', False),
        ('', False),
        ('每一行在黄色"判断"列选一项', True),
        ('对上参照点 —— 这条意见说的就是"参照点"表里的某一条。请在"参照点编号"列写上编号，例如 R03。', False),
        ('新问题·成立 —— 参照点里没有，但它确实是稿件的一个实质问题，有材料依据。', False),
        ('成立但不重要 —— 说得没错，但只是语言、格式或很细的小问题。', False),
        ('不成立 —— 说错了、误读了材料，或者推理站不住。', False),
        ('无法判断 —— 只看材料定不下来。', False),
        ('', False),
        ('怎么做', True),
        ('1. 先在"参照点"表里看这道题有哪些参照点（这是上一轮裁决后定下的参考答案）。', False),
        ('2. 打开材料：题目材料文件夹里，文件名就是题号。', False),
        ('3. 逐行判断。以材料为准，不以审稿人或参照点的说法为准。', False),
        ('4. 如果你认为参照点本身有错，请在备注里写一句。', False),
        ('5. 每道题做完，在该题第一行"本题用时"填大概分钟数。', False),
        ('', False),
        ('示例', True),
        ('意见："回应只说算过样本量，但没给效应量、α 和检验力。"　参照点 R01 是"声称计算了样本量但没有给出计算依据"。', False),
        ('→ 选"对上参照点"，参照点编号填 R01。', False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        guide.cell(row=i, column=1, value=text).font = Font(name='Arial', size=11, bold=bold)
    guide.column_dimensions['A'].width = 110

    ref = wb.create_sheet('参照点')
    for col, h in enumerate(['题号', '编号', '重要性', '内容', '怎样算对上'], 1):
        cell(ref, 1, col, h, bold=True, fill=HEAD)
    r = 2
    for task_id in chosen:
        card = json.loads((CARDS / f'{task_id}.json').read_text(encoding='utf-8'))
        for p in card['reference_points']:
            for col, v in enumerate([task_id, p['id'], IMP[p['importance']], p['judgment'], p.get('match_criteria', '')], 1):
                cell(ref, r, col, v)
            r += 1
        for p in card['neutral_points']:
            for col, v in enumerate([task_id, p['id'], '中性（不加分不扣分）', p['judgment'], p.get('match_criteria', '')], 1):
                cell(ref, r, col, v)
            r += 1
    for col, w in zip('ABCDE', (13, 7, 10, 70, 50)):
        ref.column_dimensions[col].width = w

    ws = wb.create_sheet('补评表')
    heads = ['题号', '材料文件', '答案', '模型的意见', '判断', '参照点编号', '备注', '本题用时(分)']
    for col, h in enumerate(heads, 1):
        cell(ws, 1, col, h, bold=True, fill=HEAD)
    key, row = {}, 2
    for task_id in chosen:
        task = 'T2' if '_T2_' in task_id else 'T1'
        order = list(MODELS)
        rng.shuffle(order)
        labels = dict(zip(['答案甲', '答案乙', '答案丙'], order))
        key[task_id] = {'labels': labels, 'rows': []}
        for label, model in labels.items():
            answer = json.loads((RUNS / model / f'{task_id}.json').read_text())['answer']
            for field, idx, text in claims(answer, task):
                for col, v in enumerate([task_id, f'题目材料/{task_id}.txt', label, text, '', '', '', ''], 1):
                    cell(ws, row, col, v, fill=FILL if col in (5, 6, 7, 8) else None)
                key[task_id]['rows'].append({'sheet_row': row, 'model': model, 'claim_field': field, 'claim_index': idx})
                row += 1
    for col, w in zip('ABCDEFGH', (13, 24, 8, 80, 14, 10, 30, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'E2'
    dv = DataValidation(type='list', formula1='"' + ','.join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'E2:E{row - 1}')
    wb.active = 0
    wb.save(ROOT / 'pilot/补评表_模型答案.xlsx')
    (ROOT / 'pilot/unblinded').mkdir(parents=True, exist_ok=True)
    (ROOT / 'pilot/unblinded/补评表_key.json').write_text(json.dumps({'seed': SEED, 'tasks': key}, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print('tasks', chosen, 'rows', row - 2)


if __name__ == '__main__':
    main()
