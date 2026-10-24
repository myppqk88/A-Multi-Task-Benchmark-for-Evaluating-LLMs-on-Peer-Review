#!/usr/bin/env python3
"""Apply the fixes listed in pilot/KNOWN_INPUT_ISSUES.md to the built pilot inputs, after drafting.

Edits input.json in place (the full rebuild needs the intake Release archives), records each change
in the manifest lineage and updates input_sha256 in the manifest and in evaluation/cards_pilot/.
Idempotent: rerunning on fixed inputs changes nothing.
"""
import hashlib
import json
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from build_pilot_inputs import CATALOG, clean_response_letter, dump  # noqa: E402

INPUTS = ROOT / 'pilot/inputs'
CARDS = ROOT / 'evaluation/cards_pilot'


def docx_body_order(path):
    """Paragraphs and table rows in document order (docx_text put all tables after all paragraphs)."""
    d = docx.Document(path)
    paragraphs = {p._p: p for p in d.paragraphs}
    tables = {t._tbl: t for t in d.tables}
    parts = []
    for child in d.element.body.iterchildren():
        if child in paragraphs:
            parts.append(paragraphs[child].text)
        elif child in tables:
            for row in tables[child].rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
    return '\n'.join(parts)


def fix_d17(obj, lineage, title):
    src = ROOT / lineage['responses'][0]['file']
    text = docx_body_order(src)
    text = re.sub(r'\A\s*Response to Reviewers[^\n]*\n', 'Response to Reviewers\n', text)
    text = clean_response_letter(text, title)
    text = re.sub(r'\n\s*Beverley Fredborg, Jim Clark, and Stephen Smith\s*\Z', '', text).strip()
    obj['materials']['responses']['combined'] = text
    return 'response letter re-extracted in document order (table rows in place); author-name heading and signature removed'


def fix_d08(obj, lineage, title):
    r = obj['materials']['responses']['R1']
    obj['materials']['responses']['R1'] = re.sub(r'\bWe thank Dr\.\s+\w+ for', 'We thank the reviewer for', r)
    return 'reviewer name removed from the response opening'


FIXES = {'T3_D17_R1': fix_d17, 'T3_D08_R1': fix_d08}


def main():
    manifest_path = INPUTS / 'manifest.json'
    manifest = json.loads(manifest_path.read_text())
    for item in manifest['items']:
        fix = FIXES.get(item['task_id'])
        if not fix:
            continue
        path = INPUTS / item['task_id'] / 'input.json'
        obj = json.loads(path.read_text(encoding='utf-8'))
        note = fix(obj, item['lineage'], CATALOG[item['paper_id']]['ledger']['题名'])
        dump(path, obj)
        new = hashlib.sha256(path.read_bytes()).hexdigest()
        if new == item['input_sha256']:
            print(item['task_id'], 'unchanged')
            continue
        item['lineage'].setdefault('post_draft_fixes', []).append({'note': note, 'previous_sha256': item['input_sha256']})
        item['input_sha256'] = new
        card_path = CARDS / f"{item['task_id']}.json"
        card = json.loads(card_path.read_text(encoding='utf-8'))
        card['input_sha256'] = new
        card_path.write_text(json.dumps(card, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
        print(item['task_id'], 'fixed:', note)
    dump(manifest_path, manifest)


if __name__ == '__main__':
    main()
